"""Builds the SEO page-plan workbook from plan.json + historical.json.

Usage:
  python3 .claude/skills/seo/scripts/build_plan.py <out_dir> <plan.json>
(system python3 - it has openpyxl; the project venv doesn't.)

plan.json:
{
  "file_name": "concrete-seo-page-plan-manawatu.xlsx",
  "local_label": "Palmy + Manawatu",
  "local_desc": "Palmerston North City + Manawatu District",
  "pulled": "13 Oct 2026",
  "keep_zero": ["obvious local phrasing kept for copy even at 0/0", ...],
  "pages": [{"type": "Service"|"Content", "priority": "P1"|"P2"|"P3", "page": "...",
             "url": "/...", "title": "...", "notes": "..." | null,
             "keywords": ["primary", "then in priority order", ...]}, ...],
  "excluded": [["what", "why"], ...]
}
Fails loudly if a keyword is on two pages or wasn't in the volume pull.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

out_dir, plan_path = sys.argv[1], sys.argv[2]
plan = json.load(open(plan_path))
d = json.load(open(os.path.join(out_dir, "historical.json")))
L, N = d["local"], d["nz"]
keep_zero = set(plan.get("keep_zero", []))
lab = plan["local_label"]

seen = {}
for p in plan["pages"]:
    for k in p["keywords"]:
        assert k not in seen, f"'{k}' is on both '{p['page']}' and '{seen[k]}'"
        assert k in d["kws"], f"'{k}' ({p['page']}) wasn't in the volume pull - add it to candidates and re-pull"
        seen[k] = p["page"]

def vol(src, k): return src[k] if k in src else "close variant"

wb = Workbook()
HDR = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")
def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]: c.font = HDR; c.fill = FILL; c.alignment = WRAP
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = WRAP
    ws.freeze_panes = "A2"

ws = wb.active; ws.title = "Read Me"
for line in ["How to read this",
 f"Volumes are Google Keyword Planner monthly averages, pulled {plan['pulled']}.",
 f"'{lab}' = searches made from {plan['local_desc']}. Google floors anything with any activity at 10, so most local numbers show 10 - that means 'real but small', not exactly 10.",
 f"'NZ' = nationwide searches. Use it as the tie-breaker for which keyword is more popular.",
 "'close variant' = Google merges this phrase with a near-identical one and doesn't report it separately; still worth using in copy.",
 "Keyword rank 1 = put it in the page title + H1. Ranks 2-4 = H2s / first paragraph. The rest = body copy and FAQs, written naturally (no stuffing).",
 "Every keyword is assigned to exactly one page, so pages don't compete with each other in Google.",
 "Build order: all P1 pages first, then P2, then P3. Service pages target buyers; content pages build authority and should link to the matching service page.",
 "Pages with a 'confirm with client' note depend on whether they actually offer that service - skip the page if they don't."]:
    ws.append([line])
ws["A1"].font = Font(bold=True, size=13); ws.column_dimensions["A"].width = 140

order = {"P1": 0, "P2": 1, "P3": 2}
pages = sorted(plan["pages"], key=lambda p: (p["type"] != "Service", order[p["priority"]]))
rows, kwrows = [], []
for i, p in enumerate(pages, 1):
    kws = [k for k in p["keywords"] if k in keep_zero or vol(N, k) != 0 or vol(L, k) != 0]
    nzmax = max([N[k] for k in kws if k in N] or [0])
    rows.append([i, p["type"], p["priority"], p["page"], p["url"], p["title"], kws[0], len(kws), nzmax, p.get("notes")])
    for r, k in enumerate(kws, 1):
        role = "Primary (title/H1)" if r == 1 else "Secondary (H2s / intro)" if r <= 4 else "Supporting (body / FAQs)"
        kwrows.append([i, p["page"], r, k, role, vol(L, k), vol(N, k)])

sheet(wb.create_sheet("Page Plan"), ["#", "Type", "Build priority", "Page", "Suggested URL", "Title / H1", "Primary keyword",
      "Keywords targeted", "Biggest keyword, NZ searches/mo", "Notes / confirm with client"], rows, [5, 10, 10, 40, 32, 55, 34, 11, 14, 70])
sheet(wb.create_sheet("Keywords by Page"), ["Page #", "Page", "Keyword rank", "Keyword", "Role", f"{lab} searches/mo", "NZ searches/mo"],
      kwrows, [8, 45, 9, 42, 24, 16, 14])
sheet(wb.create_sheet("Excluded & Why"), ["Excluded", "Why"], plan["excluded"], [80, 90])
path = os.path.join(out_dir, plan["file_name"])
wb.save(path)
print(f"{len(rows)} pages, {len(kwrows)} keywords -> {path}")
for r in rows: print(r[0], r[1], r[2], r[3], "|", r[6], "|", r[7], r[8])
